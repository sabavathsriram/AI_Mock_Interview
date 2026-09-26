"""
Comprehensive tests for resume upload and document processing.
"""

import pytest
import os
import tempfile
from pathlib import Path
from io import BytesIO
from datetime import datetime
from httpx import AsyncClient

from app.main import app
from app.services.resume_service import ResumeService
from app.documents import DocumentProcessorRegistry


# ============================================================================
# Fixtures for test data
# ============================================================================

@pytest.fixture
async def auth_headers(client):
    """Get authentication headers for a test user."""
    # Register a test user
    register_data = {
        "email": f"testuser_{datetime.utcnow().timestamp()}@example.com",
        "full_name": "Test User",
        "password": "TestPass123!"
    }
    response = await client.post("/api/v1/auth/register", json=register_data)
    assert response.status_code == 201
    
    # Login
    login_data = {
        "email": register_data["email"],
        "password": register_data["password"]
    }
    response = await client.post("/api/v1/auth/login", json=login_data)
    assert response.status_code == 200
    
    token_data = response.json()
    return {"Authorization": f"Bearer {token_data['access_token']}"}


@pytest.fixture
async def client():
    """Create an async test client."""
    async with AsyncClient(app=app, base_url="http://test") as c:
        yield c


# ============================================================================
# Test Sample File Generation
# ============================================================================

class TestFileGeneration:
    """Generate test files for upload testing."""
    
    @staticmethod
    def create_txt_file() -> BytesIO:
        """Create a TXT sample file."""
        content = """John Doe
Senior Software Engineer

Experience:
- 5 years as Senior Software Engineer at Tech Corp
- 3 years as Software Engineer at StartUp Inc

Skills:
- Python, Java, Go
- FastAPI, Django, Spring Boot
- MongoDB, PostgreSQL, Redis

Education:
- BS Computer Science from University of Tech"""
        
        file_obj = BytesIO(content.encode('utf-8'))
        file_obj.name = "resume.txt"
        return file_obj
    
    @staticmethod
    def create_html_file() -> BytesIO:
        """Create an HTML sample file."""
        content = """<!DOCTYPE html>
<html>
<head>
    <title>John Doe - Resume</title>
</head>
<body>
    <h1>John Doe</h1>
    <h2>Senior Software Engineer</h2>
    
    <h3>Experience</h3>
    <ul>
        <li>5 years as Senior Software Engineer at Tech Corp</li>
        <li>3 years as Software Engineer at StartUp Inc</li>
    </ul>
    
    <h3>Skills</h3>
    <p>Python, Java, Go, FastAPI, Django, Spring Boot</p>
</body>
</html>"""
        
        file_obj = BytesIO(content.encode('utf-8'))
        file_obj.name = "resume.html"
        return file_obj
    
    @staticmethod
    def create_markdown_file() -> BytesIO:
        """Create a Markdown sample file."""
        content = """# John Doe

## Senior Software Engineer

### Experience
- 5 years as Senior Software Engineer at Tech Corp
- 3 years as Software Engineer at StartUp Inc

### Skills
- Python, Java, Go
- FastAPI, Django, Spring Boot
- MongoDB, PostgreSQL, Redis

### Education
- BS Computer Science from University of Tech
"""
        
        file_obj = BytesIO(content.encode('utf-8'))
        file_obj.name = "resume.md"
        return file_obj
    
    @staticmethod
    def create_rtf_file() -> BytesIO:
        """Create an RTF sample file."""
        content = r"""{\rtf1\ansi\ansicpg1252\cocoartf2\cuc
{\fonttbl\f0\fswiss Helvetica;}
{\colortbl;\red255\green0\blue0;}
\margl1440\margr1440\margtsxn0\margbsxn0
\trowd\trgaph108\trleft-108\trbrdrt\brdrs\brdrw10 \trbrdrl\brdrs\brdrw10 
\pard\pardeftab720\sl259\slmult1\partightenfactor100

\f0\fs24 \cf0 John Doe\
Senior Software Engineer\
\
Experience:\
5 years as Senior Software Engineer at Tech Corp\
3 years as Software Engineer at StartUp Inc}"""
        
        file_obj = BytesIO(content.encode('utf-8'))
        file_obj.name = "resume.rtf"
        return file_obj
    
    @staticmethod
    def create_pdf_file() -> BytesIO:
        """Create a minimal valid PDF sample file."""
        from PyPDF2 import PdfWriter
        from io import BytesIO
        
        # Create a new PDF
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        
        # Add some basic text (PyPDF2 doesn't support adding text, so we'll create minimal PDF)
        writer.add_metadata({
            '/Title': 'John Doe Resume',
            '/Author': 'John Doe',
        })
        
        output = BytesIO()
        writer.write(output)
        output.seek(0)
        output.name = "resume.pdf"
        return output
    
    @staticmethod
    def create_docx_file() -> BytesIO:
        """Create a DOCX sample file."""
        from docx import Document
        from io import BytesIO
        
        doc = Document()
        doc.add_heading('John Doe', 0)
        doc.add_heading('Senior Software Engineer', level=1)
        
        doc.add_heading('Experience', level=2)
        doc.add_paragraph('5 years as Senior Software Engineer at Tech Corp', style='List Bullet')
        doc.add_paragraph('3 years as Software Engineer at StartUp Inc', style='List Bullet')
        
        doc.add_heading('Skills', level=2)
        doc.add_paragraph('Python, Java, Go, FastAPI, Django, Spring Boot')
        
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        output.name = "resume.docx"
        return output
    
    @staticmethod
    def create_doc_file() -> BytesIO:
        """Create a DOC sample file (using DOCX as fallback)."""
        # For testing, we'll use DOCX format and rename extension
        from docx import Document
        from io import BytesIO
        
        doc = Document()
        doc.add_heading('John Doe', 0)
        doc.add_heading('Senior Software Engineer', level=1)
        doc.add_paragraph('Experience: 5+ years in software engineering')
        
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        output.name = "resume.doc"
        return output
    
    @staticmethod
    def create_odt_file() -> BytesIO:
        """Create a minimal ODT sample file."""
        import zipfile
        from io import BytesIO
        
        # Create a minimal ODT (ZIP) file
        output = BytesIO()
        with zipfile.ZipFile(output, 'w') as zf:
            # Add mimetype
            zf.writestr('mimetype', 'application/vnd.oasis.opendocument.text')
            
            # Add minimal content.xml
            content_xml = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
                         xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">
  <office:body>
    <office:text>
      <text:p>John Doe</text:p>
      <text:p>Senior Software Engineer</text:p>
      <text:p>Experience: 5+ years</text:p>
    </office:text>
  </office:body>
</office:document-content>"""
            zf.writestr('content.xml', content_xml)
        
        output.seek(0)
        output.name = "resume.odt"
        return output
    
    @staticmethod
    def create_jpg_file() -> BytesIO:
        """Create a minimal JPG file (unsupported)."""
        # JPG magic bytes
        jpg_data = b'\xff\xd8\xff\xe0\x00\x10JFIF' + b'\x00' * 100
        file_obj = BytesIO(jpg_data)
        file_obj.name = "resume.jpg"
        return file_obj
    
    @staticmethod
    def create_png_file() -> BytesIO:
        """Create a minimal PNG file (unsupported)."""
        # PNG magic bytes
        png_data = b'\x89PNG\r\n\x1a\n' + b'\x00' * 100
        file_obj = BytesIO(png_data)
        file_obj.name = "resume.png"
        return file_obj
    
    @staticmethod
    def create_zip_file() -> BytesIO:
        """Create a ZIP file (unsupported)."""
        import zipfile
        
        output = BytesIO()
        with zipfile.ZipFile(output, 'w') as zf:
            zf.writestr('test.txt', 'This is a test')
        
        output.seek(0)
        output.name = "resume.zip"
        return output
    
    @staticmethod
    def create_exe_file() -> BytesIO:
        """Create an EXE file (unsupported)."""
        # EXE magic bytes
        exe_data = b'MZ' + b'\x00' * 100
        file_obj = BytesIO(exe_data)
        file_obj.name = "resume.exe"
        return file_obj
    
    @staticmethod
    def create_wrong_extension_file() -> BytesIO:
        """Create a TXT file with wrong extension."""
        content = "This is a text file but with wrong extension"
        file_obj = BytesIO(content.encode('utf-8'))
        file_obj.name = "resume.pdf"  # Wrong: TXT content with PDF extension
        return file_obj
    
    @staticmethod
    def create_corrupted_pdf() -> BytesIO:
        """Create a corrupted PDF file."""
        corrupted_data = b'%PDF-1.4\n' + b'\x00' * 50  # Incomplete PDF
        file_obj = BytesIO(corrupted_data)
        file_obj.name = "corrupted.pdf"
        return file_obj


# ============================================================================
# Tests for Supported File Formats
# ============================================================================

class TestSupportedFormats:
    """Test uploading all supported file formats."""
    
    @pytest.mark.asyncio
    async def test_upload_txt_file(self, client, auth_headers):
        """Test uploading a TXT file."""
        file_data = TestFileGeneration.create_txt_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.txt", file_data, "text/plain")},
            headers=auth_headers
        )
        
        assert response.status_code == 201
        data = response.json()
        assert data["file_type"] == "txt"
        assert data["filename"] == "resume.txt"
        assert len(data["extracted_text"]) > 0
        assert "John Doe" in data["extracted_text"]
    
    @pytest.mark.asyncio
    async def test_upload_html_file(self, client, auth_headers):
        """Test uploading an HTML file."""
        file_data = TestFileGeneration.create_html_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.html", file_data, "text/html")},
            headers=auth_headers
        )
        
        assert response.status_code == 201
        data = response.json()
        assert data["file_type"] == "html"
        assert len(data["extracted_text"]) > 0
        assert "John Doe" in data["extracted_text"]
    
    @pytest.mark.asyncio
    async def test_upload_markdown_file(self, client, auth_headers):
        """Test uploading a Markdown file."""
        file_data = TestFileGeneration.create_markdown_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.md", file_data, "text/markdown")},
            headers=auth_headers
        )
        
        assert response.status_code == 201
        data = response.json()
        assert data["file_type"] == "markdown"
        assert len(data["extracted_text"]) > 0
    
    @pytest.mark.asyncio
    async def test_upload_rtf_file(self, client, auth_headers):
        """Test uploading an RTF file."""
        file_data = TestFileGeneration.create_rtf_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.rtf", file_data, "application/rtf")},
            headers=auth_headers
        )
        
        assert response.status_code == 201
        data = response.json()
        assert data["file_type"] == "rtf"
        assert len(data["extracted_text"]) > 0
    
    @pytest.mark.asyncio
    async def test_upload_pdf_file(self, client, auth_headers):
        """Test uploading a PDF file."""
        file_data = TestFileGeneration.create_pdf_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.pdf", file_data, "application/pdf")},
            headers=auth_headers
        )
        
        assert response.status_code == 201
        data = response.json()
        assert data["file_type"] == "pdf"
    
    @pytest.mark.asyncio
    async def test_upload_docx_file(self, client, auth_headers):
        """Test uploading a DOCX file."""
        file_data = TestFileGeneration.create_docx_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.docx", file_data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=auth_headers
        )
        
        assert response.status_code == 201
        data = response.json()
        assert data["file_type"] == "docx"
        assert "John Doe" in data["extracted_text"]
    
    @pytest.mark.asyncio
    async def test_upload_doc_file(self, client, auth_headers):
        """Test uploading a DOC file."""
        file_data = TestFileGeneration.create_doc_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.doc", file_data, "application/msword")},
            headers=auth_headers
        )
        
        assert response.status_code == 201
        data = response.json()
        assert data["file_type"] == "doc"
    
    @pytest.mark.asyncio
    async def test_upload_odt_file(self, client, auth_headers):
        """Test uploading an ODT file."""
        file_data = TestFileGeneration.create_odt_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.odt", file_data, "application/vnd.oasis.opendocument.text")},
            headers=auth_headers
        )
        
        assert response.status_code == 201
        data = response.json()
        assert data["file_type"] == "odt"


# ============================================================================
# Tests for Unsupported Formats
# ============================================================================

class TestUnsupportedFormats:
    """Test rejection of unsupported file formats."""
    
    @pytest.mark.asyncio
    async def test_reject_jpg_file(self, client, auth_headers):
        """Test that JPG files are rejected."""
        file_data = TestFileGeneration.create_jpg_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.jpg", file_data, "image/jpeg")},
            headers=auth_headers
        )
        
        assert response.status_code == 400
        data = response.json()
        assert "JPG" in data["detail"] or "image" in data["detail"].lower()
    
    @pytest.mark.asyncio
    async def test_reject_png_file(self, client, auth_headers):
        """Test that PNG files are rejected."""
        file_data = TestFileGeneration.create_png_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.png", file_data, "image/png")},
            headers=auth_headers
        )
        
        assert response.status_code == 400
        data = response.json()
        assert "PNG" in data["detail"] or "image" in data["detail"].lower()
    
    @pytest.mark.asyncio
    async def test_reject_zip_file(self, client, auth_headers):
        """Test that ZIP files are rejected."""
        file_data = TestFileGeneration.create_zip_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.zip", file_data, "application/zip")},
            headers=auth_headers
        )
        
        assert response.status_code == 400
        data = response.json()
        assert "ZIP" in data["detail"] or "not supported" in data["detail"].lower()
    
    @pytest.mark.asyncio
    async def test_reject_exe_file(self, client, auth_headers):
        """Test that EXE files are rejected."""
        file_data = TestFileGeneration.create_exe_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.exe", file_data, "application/x-msdownload")},
            headers=auth_headers
        )
        
        assert response.status_code == 400
        data = response.json()
        assert "Executable" in data["detail"] or "not supported" in data["detail"].lower()


# ============================================================================
# Tests for File Validation
# ============================================================================

class TestFileValidation:
    """Test file validation logic."""
    
    @pytest.mark.asyncio
    async def test_reject_wrong_extension(self, client, auth_headers):
        """Test that files with wrong extensions are rejected."""
        file_data = TestFileGeneration.create_wrong_extension_file()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.pdf", file_data, "text/plain")},
            headers=auth_headers
        )
        
        assert response.status_code == 400
        data = response.json()
        assert "corrupted" in data["detail"].lower() or "invalid" in data["detail"].lower()
    
    @pytest.mark.asyncio
    async def test_reject_corrupted_pdf(self, client, auth_headers):
        """Test that corrupted PDFs are rejected."""
        file_data = TestFileGeneration.create_corrupted_pdf()
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("corrupted.pdf", file_data, "application/pdf")},
            headers=auth_headers
        )
        
        assert response.status_code == 400
        data = response.json()
        assert "extract" in data["detail"].lower() or "failed" in data["detail"].lower()
    
    @pytest.mark.asyncio
    async def test_file_size_validation(self, client, auth_headers):
        """Test that oversized files are rejected."""
        # Create a file larger than 10MB
        large_content = b'A' * (11 * 1024 * 1024)  # 11MB
        file_data = BytesIO(large_content)
        file_data.name = "large_resume.txt"
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("large_resume.txt", file_data, "text/plain")},
            headers=auth_headers
        )
        
        assert response.status_code == 413  # Payload Too Large
        data = response.json()
        assert "exceeds" in data["detail"].lower() or "size" in data["detail"].lower()
    
    @pytest.mark.asyncio
    async def test_empty_file_rejection(self, client, auth_headers):
        """Test that empty files are rejected."""
        file_data = BytesIO(b'')
        
        response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("empty.txt", file_data, "text/plain")},
            headers=auth_headers
        )
        
        assert response.status_code == 400
        data = response.json()
        assert "empty" in data["detail"].lower() or "size" in data["detail"].lower()


# ============================================================================
# Tests for Resume Management
# ============================================================================

class TestResumeManagement:
    """Test resume listing, retrieval, and deletion."""
    
    @pytest.mark.asyncio
    async def test_list_user_resumes(self, client, auth_headers):
        """Test listing user's resumes."""
        # Upload a resume first
        file_data = TestFileGeneration.create_txt_file()
        upload_response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.txt", file_data, "text/plain")},
            headers=auth_headers
        )
        assert upload_response.status_code == 201
        
        # List resumes
        response = await client.get(
            "/api/v1/resumes/list",
            headers=auth_headers
        )
        
        assert response.status_code == 200
        data = response.json()
        assert "resumes" in data
        assert "total_count" in data
        assert data["total_count"] >= 1
        assert len(data["resumes"]) >= 1
        assert data["resumes"][0]["filename"] == "resume.txt"
    
    @pytest.mark.asyncio
    async def test_get_resume_details(self, client, auth_headers):
        """Test getting resume details."""
        # Upload a resume
        file_data = TestFileGeneration.create_txt_file()
        upload_response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.txt", file_data, "text/plain")},
            headers=auth_headers,
            data={"display_name": "My First Resume"}
        )
        assert upload_response.status_code == 201
        resume_id = upload_response.json()["id"]
        
        # Get resume details
        response = await client.get(
            f"/api/v1/resumes/{resume_id}",
            headers=auth_headers
        )
        
        assert response.status_code == 200
        data = response.json()
        assert data["id"] == resume_id
        assert data["filename"] == "resume.txt"
        assert data["file_type"] == "txt"
        assert data["extraction_status"] == "success"
    
    @pytest.mark.asyncio
    async def test_delete_resume(self, client, auth_headers):
        """Test deleting a resume."""
        # Upload a resume
        file_data = TestFileGeneration.create_txt_file()
        upload_response = await client.post(
            "/api/v1/resumes/upload",
            files={"file": ("resume.txt", file_data, "text/plain")},
            headers=auth_headers
        )
        assert upload_response.status_code == 201
        resume_id = upload_response.json()["id"]
        
        # Delete resume
        response = await client.delete(
            f"/api/v1/resumes/{resume_id}",
            headers=auth_headers
        )
        
        assert response.status_code == 200
        data = response.json()
        assert data["success"] is True
        
        # Verify it's deleted
        get_response = await client.get(
            f"/api/v1/resumes/{resume_id}",
            headers=auth_headers
        )
        assert get_response.status_code == 404


# ============================================================================
# Tests for Supported Formats Information
# ============================================================================

class TestFormatsInfo:
    """Test the supported formats information endpoint."""
    
    @pytest.mark.asyncio
    async def test_get_supported_formats(self, client):
        """Test getting supported formats information."""
        response = await client.get("/api/v1/resumes/info/supported-formats")
        
        assert response.status_code == 200
        data = response.json()
        assert "supported_formats" in data
        assert "unsupported_formats" in data
        assert "max_file_size_mb" in data
        assert "max_file_size_bytes" in data
        
        # Verify all supported formats are listed
        supported = data["supported_formats"]
        assert ".txt" in supported
        assert ".pdf" in supported
        assert ".docx" in supported
        assert ".doc" in supported
        assert ".rtf" in supported
        assert ".odt" in supported
        assert ".html" in supported
        assert ".md" in supported
        
        # Verify unsupported formats are listed
        unsupported = data["unsupported_formats"]
        assert ".jpg" in unsupported or ".jpeg" in unsupported
        assert ".png" in unsupported
        assert ".zip" in unsupported
        assert ".exe" in unsupported
