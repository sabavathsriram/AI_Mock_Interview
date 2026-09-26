"""
DOCX document processor using python-docx.
"""

from typing import Dict, Any
from app.documents.base import DocumentProcessor, ExtractionResult
import os


class DOCXProcessor(DocumentProcessor):
    """Processor for DOCX documents."""
    
    supported_extensions = ('.docx',)
    mime_types = ('application/vnd.openxmlformats-officedocument.wordprocessingml.document',)
    
    async def extract(self, file_path: str) -> ExtractionResult:
        """
        Extract text from DOCX file.
        
        Args:
            file_path: Path to the DOCX file
            
        Returns:
            ExtractionResult containing extracted text and metadata
        """
        try:
            from docx import Document
            
            if not os.path.exists(file_path):
                return ExtractionResult(
                    text="",
                    metadata={},
                    file_type="docx",
                    success=False,
                    error="File not found"
                )
            
            try:
                doc = Document(file_path)
            except Exception as e:
                return ExtractionResult(
                    text="",
                    metadata={},
                    file_type="docx",
                    success=False,
                    error=f"Failed to read DOCX: {str(e)}"
                )
            
            text = ""
            paragraph_count = 0
            table_count = 0
            
            # Extract text from paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text += para.text + "\n"
                    paragraph_count += 1
            
            # Extract text from tables
            for table in doc.tables:
                table_count += 1
                text += "\n[TABLE]\n"
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    if row_text.strip():
                        text += row_text + "\n"
            
            # Normalize the text
            normalized_text = self.normalize_text(text)
            
            metadata: Dict[str, Any] = {
                "paragraph_count": paragraph_count,
                "table_count": table_count,
                "processor": "python-docx"
            }
            
            return ExtractionResult(
                text=normalized_text,
                metadata=metadata,
                file_type="docx",
                success=True
            )
        
        except Exception as e:
            return ExtractionResult(
                text="",
                metadata={},
                file_type="docx",
                success=False,
                error=f"DOCX extraction failed: {str(e)}"
            )
